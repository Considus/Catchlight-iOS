#!/usr/bin/env python3
"""Build a translator review workbook (.xlsx) for one language from the String Catalogs.

Every string in English beside the draft, a yellow column for the translator's
version, per-row English word counts, a glossary and a read-me with the totals.
Each language's notes, glossary and tone live in LANGS below; add one there to
support a new language.

--store adds the App Store listing and screenshot text, which live outside the
repo: a source file of rows {id, where, form, en, limit, note, en_us} and the
language's draft {id: text}.

--texts DIR writes plain-text files instead: Catchlight_iOS_Translation_EN.txt
and one Catchlight_iOS_Translation_<LANG>.txt per language, where line N of each
translates line N of the English, in plain language (see readable()). With
--store-src SRC --store-dir DIR they include the App Store rows, read from
DIR/store-<lang>.json.

Usage: review_workbook.py <lang> <out.xlsx> [--root DIR] [--store SRC DRAFT]
       review_workbook.py --texts DIR [--root DIR] [--store-src SRC --store-dir DIR]
Needs openpyxl.
"""
import argparse
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

AREA = {
    "Catchlight/Resources/Localizable.xcstrings": "App",
    "Catchlight/Resources/InfoPlist.xcstrings": "Permission prompt",
    "Catchlight/Resources/AppShortcuts.xcstrings": "Siri phrase",
    "CatchlightWidgets/Localizable.xcstrings": "Widget",
    "CatchlightShare/Localizable.xcstrings": "Share sheet",
}

KEEP = "Take, Obie, Iris, Dailies, Sequence, Angle, Shot List, Storyboard, Catchlight, Considus"

LANGS = {
    "de": {
        "name": "German",
        "tone": 'Informal "du", warm and plain. Catchlight is a private notes app.',
        "plural": 'Rows marked "plural: one" are used when the count is 1; "plural: other" for every other count.',
        "example": 'English "Couldn\'t load your Takes." · draft "Deine Takes konnten nicht geladen werden." · '
                   'your German "Deine Takes ließen sich nicht laden." · your comment "shorter, more natural"',
        "notes": {
            "Privacy phrase": 'Glossary term. Draft "Datenschutz-Phrase": please confirm or propose a better term; it recurs in many strings.',
            "Opens the Focus ring to change what this Take is.": '"Focus ring" drafted as "Fokusring". Avoid confusion with the iOS feature "Fokus".',
            "%@ at %@": 'An email address read aloud by VoiceOver (user "at" domain). Draft keeps "at"; change if German VoiceOver users expect something else.',
            "%lld-day": 'Adjective inside "Start your %@ trial now…" (e.g. "14-tägige Testphase"). Must agree with the noun in that sentence.',
            "Start your %@ trial now, then only %@": 'First %@ is the trial-length adjective (e.g. "14-tägige"), second is the price (e.g. "14,99 €/Jahr").',
            "Not configured": 'Shown in a narrow row; drafted as "Kein Ordner" so it fits.',
            "Snooze for %@": '%@ is a duration such as "1 Stunde".',
        },
        "glossary": [
            ("Take", "Take (der Take, die Takes)", "Product name. Keep in English. Masculine."),
            ("Obie", "Obie (das Obie)", "Product name. Keep in English. Neuter."),
            ("Iris", "Iris (die Iris)", "Product name: the circle beside each Take. Keep."),
            ("Dailies", "Dailies (die Dailies)", "Product name: the main timeline. Keep."),
            ("Sequence", "Sequence (die Sequence)", "Product name: a saved filter. Keep."),
            ("Angle", "Angle (der Angle)", "Product name. Keep."),
            ("Shot List", "Shot List (die Shot List)", "Product name: a Take's checklist view. Keep."),
            ("Storyboard", "Storyboard (das Storyboard)", "Product name: every Take with a task. Keep."),
            ("Privacy phrase", "Datenschutz-Phrase", "The 12 recovery words. Draft term, please confirm."),
            ("Focus ring", "Fokusring", "Draft term, please confirm."),
            ("Note / Task / Reminder", "Notiz / Aufgabe / Erinnerung", "Ordinary words, translated."),
            ("Save", "Sichern", "Follows Apple's German UI."),
            ("Double-tap …", "Doppeltippen …", "VoiceOver hints follow Apple's German VoiceOver wording."),
            ("you", "du / dein", "Informal throughout, lower case."),
        ],
    },
    "fr": {
        "name": "French",
        "tone": 'Formal "vous", warm and plain. Catchlight is a private notes app. French typography: a non-breaking '
                'space before : and a narrow one before ? ! ;, « guillemets », curly apostrophes.',
        "plural": 'Rows marked "plural: one" are used when the count is 0 or 1; "plural: other" for every other count.',
        "example": 'English "Couldn\'t load your Takes." · draft "Impossible de charger vos Takes." · '
                   'your French "Vos Takes n\'ont pas pu être chargés." · your comment "softer"',
        "notes": {
            "Privacy phrase": 'Glossary term. Draft "phrase de confidentialité": please confirm or propose a better term; it recurs in many strings.',
            "Opens the Focus ring to change what this Take is.": '"Focus ring" drafted as "anneau de focus". The iOS feature "Focus" is "mode de concentration" in French and appears separately.',
            "Which Takes Catchlight shows while this Focus is on.": 'Here "Focus" is the iOS feature, which Apple calls "mode de concentration".',
            "%@ at %@": 'An email address read aloud by VoiceOver (user "at" domain). Drafted with "arobase".',
            "%lld-day": 'Trial length inside "Start your %@ trial now…" (e.g. "essai de 14 jours"). French needs a plural here; English does not.',
            "Start your %@ trial now, then only %@": 'First %@ is the trial length (e.g. "14 jours"), second is the price (e.g. "14,99 €/an").',
            "Snooze for %@": '%@ is a duration such as "1 heure".',
            "Settings": 'Drafted as "Réglages", the name Apple uses for Settings in French iOS, for the app\'s own screen as well.',
            "Notice History": '"Notice" (a status message on the main screen) drafted as "alerte" throughout.',
            "Snooze": 'Drafted as "Reporter" rather than Apple\'s alarm wording "Répéter", which here already means Repeat.',
        },
        "glossary": [
            ("Take", "Take (le Take, les Takes)", "Product name. Keep in English. Masculine."),
            ("Obie", "Obie (l'Obie, un Obie)", "Product name. Keep in English. Masculine: \"Nouvel Obie\"."),
            ("Iris", "Iris (l'Iris, un Iris)", "Product name: the circle beside each Take. Keep. Masculine, as the French noun."),
            ("Dailies", "Dailies", "Product name: the main timeline. Keep."),
            ("Sequence", "Sequence", "Product name: a saved filter. Keep."),
            ("Angle", "Angle", "Product name. Keep."),
            ("Shot List", "Shot List (la Shot List)", "Product name: a Take's checklist view. Keep."),
            ("Storyboard", "Storyboard (le Storyboard)", "Product name: every Take with a task. Keep."),
            ("Privacy phrase", "phrase de confidentialité", "The 12 recovery words. Draft term, please confirm."),
            ("Focus ring", "anneau de focus", "Draft term, please confirm."),
            ("Timeline", "chronologie", "Draft term."),
            ("Notice", "alerte", "A status message on the main screen. Draft term."),
            ("Note / Task / Reminder", "note / tâche / rappel", "Ordinary words, translated."),
            ("Settings", "Réglages", "Follows Apple's French UI."),
            ("Save", "Enregistrer", "Follows Apple's French UI."),
            ("Double-tap …", "Touchez deux fois …", "VoiceOver hints follow Apple's French VoiceOver wording."),
            ("you", "vous / votre", "Formal throughout."),
        ],
    },
    "es": {
        "name": "Spanish",
        "tone": 'Informal "tú", warm and plain. Catchlight is a private notes app. Spanish as used in Spain.',
        "plural": 'Rows marked "plural: one" are used when the count is 1; "plural: other" for every other count.',
        "example": 'English "Couldn\'t load your Takes." · draft "No se han podido cargar tus Takes." · '
                   'your Spanish "No se pudieron cargar tus Takes." · your comment "más natural"',
        "notes": {
            "Privacy phrase": 'Glossary term. Draft "frase de privacidad": please confirm or propose a better term.',
            "%lld-day": 'Trial length inside "Start your %@ trial now…" (e.g. "prueba de 14 días"). Spanish needs a plural here; English does not.',
            "Done": 'One string serves as a button and as a task state; drafted "Hecho".',
            "Created on %@ at %@": 'Drafted with "a las", which is wrong at 1 o\'clock ("a la 1"). A wording that works for every hour is welcome.',
        },
        "glossary": [
            ("Take", "Take (el Take, los Takes)", "Product name. Keep in English. Masculine."),
            ("Obie", "Obie (el Obie)", "Product name. Keep in English. Masculine."),
            ("Iris", "Iris (el Iris)", "Product name: the circle beside each Take. Keep. Masculine."),
            ("Dailies, Sequence, Angle, Shot List, Storyboard", "unchanged", "Product names. Keep."),
            ("Privacy phrase", "frase de privacidad", "The 12 recovery words. Draft term, please confirm."),
            ("Timeline", "cronología", "Draft term."),
            ("Notice", "aviso", "A status message on the main screen. Draft term."),
            ("Settings / Save / Snooze", "Ajustes / Guardar / Posponer", "Follows Apple's Spanish UI."),
            ("Double-tap …", "Toca dos veces …", "VoiceOver hints follow Apple's Spanish wording."),
            ("you", "tú / tu", "Informal throughout."),
        ],
    },
    "it": {
        "name": "Italian",
        "tone": 'Informal "tu", warm and plain. Catchlight is a private notes app.',
        "plural": 'Rows marked "plural: one" are used when the count is 1; "plural: other" for every other count.',
        "example": 'English "Couldn\'t load your Takes." · draft "Impossibile caricare i tuoi Take." · '
                   'your Italian "Non è stato possibile caricare i tuoi Take." · your comment "meno brusco"',
        "notes": {
            "Privacy phrase": 'Glossary term. Draft "frase di privacy": please confirm or propose a better term.',
            "%lld-day": 'Trial length inside "Start your %@ trial now…" (e.g. "prova di 14 giorni"). Italian needs a plural here; English does not.',
            "Done": 'One string serves as a button and as a task state.',
            "Inline": "Apple's Italian name for this Writing Tools option was not confirmed.",
        },
        "glossary": [
            ("Take", "Take (il Take, i Take)", "Product name. Keep in English. Masculine, invariable plural."),
            ("Obie", "Obie (l'Obie, gli Obie)", "Product name. Keep in English. Masculine."),
            ("Iris", "Iris (l'Iris)", "Product name: the circle beside each Take. Keep. Masculine."),
            ("Dailies, Sequence, Angle, Shot List, Storyboard", "unchanged", "Product names. Keep."),
            ("Privacy phrase", "frase di privacy", "The 12 recovery words. Draft term, please confirm."),
            ("Timeline", "cronologia", "Draft term."),
            ("Notice", "avviso", "A status message on the main screen. Draft term."),
            ("Settings / Save / Snooze", "Impostazioni / Salva / Posticipa", "Follows Apple's Italian UI."),
            ("Double-tap …", "Tocca due volte …", "VoiceOver hints follow Apple's Italian wording."),
            ("you", "tu / tuo", "Informal throughout."),
        ],
    },
    "nl": {
        "name": "Dutch",
        "tone": 'Informal "je / jouw", as Apple uses. Warm and plain. Catchlight is a private notes app.',
        "plural": 'Rows marked "plural: one" are used when the count is 1; "plural: other" for every other count.',
        "example": 'English "Couldn\'t load your Takes." · draft "Je Takes konden niet worden geladen." · '
                   'your Dutch "Kan je Takes niet laden." · your comment "korter"',
        "notes": {
            "Privacy phrase": 'Glossary term. Draft "privacyzin": please confirm or propose a better term.',
            "Notice History": '"Notice" is drafted as "melding", but Apple\'s Dutch for Notifications is also "Meldingen". A clearer word for these status messages is welcome.',
            "%lld-day": 'Trial length inside "Start your %@ trial now…" (e.g. "proefperiode van 14 dagen"). Dutch needs a plural here; English does not.',
            "Done": 'One string serves as a button and as a task state; drafted "Klaar".',
        },
        "glossary": [
            ("Take", "Take (de Take, de Takes)", "Product name. Keep in English."),
            ("Obie", "Obie (de Obie)", "Product name. Keep in English."),
            ("Iris", "Iris (de Iris)", "Product name: the circle beside each Take. Keep."),
            ("Dailies, Sequence, Angle, Shot List, Storyboard", "unchanged", "Product names. Keep."),
            ("Privacy phrase", "privacyzin", "The 12 recovery words. Draft term, please confirm."),
            ("Timeline", "tijdlijn", "Draft term."),
            ("Notice", "melding", "A status message on the main screen. Draft term; clashes with Notifications."),
            ("Settings / Save / Snooze", "Instellingen / Bewaar / Sluimer", "Follows Apple's Dutch UI."),
            ("Double-tap …", "Tik twee keer …", "VoiceOver hints follow Apple's Dutch wording."),
            ("you", "je / jouw", "Informal throughout."),
        ],
    },
    "pt-BR": {
        "name": "Brazilian Portuguese",
        "tone": 'Informal "você", warm and plain. Catchlight is a private notes app. Brazilian usage.',
        "plural": 'Rows marked "plural: one" are used when the count is 0 or 1; "plural: other" for every other count.',
        "example": 'English "Couldn\'t load your Takes." · draft "Não foi possível carregar seus Takes." · '
                   'your Portuguese "Não deu para carregar seus Takes." · your comment "mais leve"',
        "notes": {
            "Privacy phrase": 'Glossary term. Draft "frase de privacidade": please confirm or propose a better term.',
            "%lld-day": 'Trial length inside "Start your %@ trial now…" (e.g. "teste de 14 dias"). Portuguese needs a plural here; English does not.',
            "Done": 'One string serves as a button and as a task state; drafted "Concluído".',
            "Welcome back": 'Drafted gender-neutral ("Que bom ter você de volta"); a shorter neutral form is welcome.',
        },
        "glossary": [
            ("Take", "Take (o Take, os Takes)", "Product name. Keep in English. Masculine."),
            ("Obie", "Obie (o Obie)", "Product name. Keep in English. Masculine."),
            ("Iris", "Iris (o Iris)", "Product name: the circle beside each Take. Keep. Masculine."),
            ("Dailies, Sequence, Angle, Shot List, Storyboard", "unchanged", "Product names. Keep."),
            ("Privacy phrase", "frase de privacidade", "The 12 recovery words. Draft term, please confirm."),
            ("Timeline", "linha do tempo", "Draft term."),
            ("Notice", "aviso", "A status message on the main screen. Draft term."),
            ("Settings / Save / Snooze", "Ajustes / Salvar / Adiar", "Follows Apple's Brazilian Portuguese UI."),
            ("Double-tap …", "Toque duas vezes …", "VoiceOver hints follow Apple's Brazilian Portuguese wording."),
            ("you", "você / seu", "Informal throughout."),
        ],
    },
}


def collect(root, lang):
    rows = []
    for rel in sorted(AREA):
        d = json.load(open(os.path.join(root, rel)))
        for k, e in d["strings"].items():
            loc = e.get("localizations", {})
            tr, en = loc.get(lang), loc.get("en")
            if not tr:
                continue
            comment = e.get("comment", "")
            en_unit = en["stringUnit"]["value"] if en and "stringUnit" in en else k
            if "stringSet" in tr:
                for i, (ev, tv) in enumerate(zip(en["stringSet"]["values"], tr["stringSet"]["values"])):
                    rows.append((AREA[rel], k, f"phrase {i + 1}", ev, tv,
                                 comment or "Something a person says to Siri. Must contain ${applicationName}."))
            elif "variations" in tr:
                ep = en["variations"]["plural"] if en and "variations" in en else None
                for form, sub in tr["variations"]["plural"].items():
                    ev = ep[form]["stringUnit"]["value"] if ep and form in ep else en_unit
                    rows.append((AREA[rel], k, f"plural: {form}", ev, sub["stringUnit"]["value"], comment))
            else:
                if k.startswith("CFBundle") or not re.search(r"[A-Za-z]", re.sub(r"%(\d\$)?(lld|@)", "", en_unit)):
                    continue
                rows.append((AREA[rel], k, "", en_unit, tr["stringUnit"]["value"], comment))
    return rows


def store_rows(src_path, draft_path):
    src = json.load(open(src_path, encoding="utf-8"))
    draft = json.load(open(draft_path, encoding="utf-8"))
    missing = [r["id"] for r in src if r["id"] not in draft]
    if missing:
        sys.exit(f"store draft is missing: {', '.join(missing)}")
    rows = []
    for r in src:
        note = r.get("note", "")
        if r.get("limit"):
            note = f"Max {r['limit']} characters. {note}".strip()
        if r.get("en_us"):
            note = f"{note} U.S. English reads: \"{r['en_us']}\"".strip()
        rows.append((r["where"], r["id"], r["form"], r["en"], draft[r["id"]], note))
    return rows


TOKENS = {"${applicationName}": "{app}", "${text}": "{text}", "${scope}": "{filter}"}
PRODUCT = re.compile(r"\b(Takes?|Obies?|Iris|Dailies|DAILIES|Sequence|SEQUENCE|Shot List|SHOT LIST|"
                     r"Storyboard|STORYBOARD|Angle|Catchlight|Considus)\b")
SPEC = re.compile(r"%(?:(\d+)\$)?(?:l{0,2}[@dDuUxXoOfeEgGcCsSp])")


def readable(text):
    """A string as a translator reads it: placeholders become {1}, {2}… by position
    (so a translation may reorder them), Siri tokens become {app}, {text}, {filter},
    a line break inside a string becomes " / " and bold markers are dropped. It is
    one-way: the text files are for reading and translating; corrections go back into
    the catalogs through the workbook or by hand, restoring any bold markers and line
    breaks from the English key."""
    auto = 0

    def number(m):
        nonlocal auto
        if m.group(1):
            return "{" + m.group(1) + "}"
        auto += 1
        return "{" + str(auto) + "}"

    text = SPEC.sub(number, text)
    for raw, shown in TOKENS.items():
        text = text.replace(raw, shown)
    return text.replace("\n", " / ").replace("**", "")


def english_line(row):
    """The English a translator reads for one row. A count with no English plural
    ("%lld-day", the trial length) reads as "{1} day" / "{1} days" so its two forms differ."""
    _, key, form, en, _, _ = row
    m = re.fullmatch(r"%lld-(\w+)", key)
    if m and form.startswith("plural:"):
        return "{1} " + m.group(1) + ("" if form == "plural: one" else "s")
    return readable(str(en))


def write_texts(out_dir, root, store_src, store_dir):
    """One English file and one file per language in LANGS: line N of each language
    file translates line N of the English. The line list is shared, so it is the union
    of every language's rows (a language with no plural for a key repeats its one
    translation on both lines), each English line once, and only lines with something
    to translate. Rows that share an English line must share its translation, so one
    line stands for every key whose English reads as that line."""
    per_lang = {}
    for lang in LANGS:
        rows = collect(root, lang)
        if store_src:
            rows += store_rows(store_src, os.path.join(store_dir, f"store-{lang}.json"))
        per_lang[lang] = {(r[0], r[1], r[2]): r for r in rows}
    split = {i[:2] for rows in per_lang.values() for i in rows if i[2].startswith("plural:")}
    order, english = [], {}
    for lang, rows in per_lang.items():
        for ident, r in rows.items():
            if ident in english or (ident[2] == "" and ident[:2] in split):
                continue
            en = english_line(r)
            rest = PRODUCT.sub("", re.sub(r"\{\w+\}", "", en))
            if re.search(r"[A-Za-z]{2,}", rest):
                english[ident] = en
                order.append(ident)
    def find(rows, ident):
        if ident in rows:
            return rows[ident]
        area, key, _ = ident  # a plural form this language does not split
        return next((r for i, r in rows.items() if i[:2] == (area, key)), None)

    # Every kept row must be translated, and rows sharing an English line must share
    # its translation in every language, or the merged line would hide a difference.
    # Checked before anything is written, so a failure leaves the old files whole.
    text, problems = {}, []
    for lang, rows in per_lang.items():
        by_line = {}
        for ident in order:
            r = find(rows, ident)
            if r is None:
                problems.append(f"{lang}: no translation for {ident}")
                continue
            by_line.setdefault(english[ident], set()).add(readable(str(r[4])))
        for en, versions in by_line.items():
            if len(versions) > 1:
                problems.append(f"{lang}: {en!r} is translated {len(versions)} ways: {sorted(versions)}")
        text[lang] = {en: next(iter(v)) for en, v in by_line.items()}
    if problems:
        sys.exit("text files not written:\n  " + "\n  ".join(problems))
    lines = list(dict.fromkeys(english[i] for i in order))

    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, "Catchlight_iOS_Translation_EN.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    for lang in per_lang:
        path = os.path.join(out_dir, f"Catchlight_iOS_Translation_{lang.upper()}.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(text[lang][en] for en in lines) + "\n")
    print(f"{out_dir}: EN + {len(per_lang)} languages, {len(lines)} lines each")


def words(s):
    s = re.sub(r"%(\d\$)?(lld|ld|d|@)|\$\{\w+\}|\*\*", " ", s)
    return len([w for w in s.split() if re.search(r"[A-Za-zÀ-ÿ]", w)])


def build(lang, out, root, store=None):
    cfg = LANGS[lang]
    name = cfg["name"]
    rows = collect(root, lang)
    if store:
        rows += store_rows(*store)

    wb = Workbook()
    F = "Arial"
    hdr_font = Font(name=F, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="5B4A2F")
    body = Font(name=F, size=10)
    wrap = Alignment(wrap_text=True, vertical="top")
    fill_in = PatternFill("solid", fgColor="FFF2B3")
    thin = Border(bottom=Side(style="thin", color="DDDDDD"))

    ws = wb.active
    ws.title = "Strings"
    heads = ["ID", "Where", "Form", "English (source)", f"{name} (draft)", "Context", "Notes for the translator",
             f"Your {name} (only if different)", "Your comment", "English words"]
    widths = [6, 14, 12, 48, 48, 30, 36, 44, 30, 9]
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment = hdr_font, hdr_fill, wrap
        ws.column_dimensions[cell.column_letter].width = w
    total = 0
    for i, (area, key, form, ev, tv, comment) in enumerate(rows, 2):
        n = words(ev)
        total += n
        for c, v in enumerate([i - 1, area, form, ev, tv, comment, cfg["notes"].get(key, ""), None, None, n], 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font, cell.alignment, cell.border = body, wrap, thin
        ws.cell(row=i, column=8).fill = fill_in
        ws.cell(row=i, column=9).fill = fill_in
    last = len(rows) + 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:J{last}"

    g = wb.create_sheet("Glossary")
    gl = [("Term (English)", name, "Rule")] + cfg["glossary"] + [("Catchlight, Considus", "unchanged", "Brand names. Keep.")]
    for r, row in enumerate(gl, 1):
        for c, v in enumerate(row, 1):
            cell = g.cell(row=r, column=c, value=v)
            cell.alignment = wrap
            cell.font = hdr_font if r == 1 else body
            if r == 1:
                cell.fill = hdr_fill
    for col, w in zip("ABC", (24, 32, 60)):
        g.column_dimensions[col].width = w

    rm = wb.create_sheet("Read me", 0)
    lines = [
        (f"Catchlight for iOS: {name} translation review", None),
        ("", None),
        ("What this is", "Every piece of text in the Catchlight iPhone app (screens, VoiceOver, notifications, widgets, "
                         f"share sheet, Siri phrases, permission prompts) in English, with a {name} draft beside it."),
        ("What we need", f"Read each {name} draft. If it is right, leave it. If not, write your {name} in the yellow "
                         f'column "Your {name} (only if different)" on the Strings tab, and add a comment if useful.'),
        ("Total English words", f"=SUM(Strings!J2:J{last})"),
        ("Strings to review", f"=COUNTA(Strings!A2:A{last})"),
        ("Tone", cfg["tone"]),
        ("Keep in English", f"Product names: {KEEP}. Their genders are on the Glossary tab."),
        ("Placeholders", "Keep every %@, %lld, %1$@, ${text}, ${scope} and ${applicationName} exactly as written. "
                         "The app replaces them with names, numbers or dates. Numbered ones (%1$@, %2$@) may move "
                         "within the sentence."),
        ("Plurals", cfg["plural"]),
        ("Siri phrases", 'Rows marked "phrase N" are alternative things a person can say to Siri. Each must contain '
                         "${applicationName}."),
        ("Length", f"Short labels sit in narrow rows on an iPhone screen. Where the {name} runs much longer than the "
                   "English, a shorter wording is welcome."),
        ("App Store and screenshots", 'Rows whose "Where" is App Store or Screenshot are the store listing and the text '
                                      "shown in the App Store screenshots. Character limits are in the Context column and "
                                      "are hard limits. Keywords are search terms people in your country would type, not "
                                      "a translation. Demo Takes are sample notes: make them feel local."),
        ("Word count", "English words per row are in the last column of the Strings tab: words containing letters, "
                       "placeholders not counted. Each plural form and each Siri phrase counts as its own row."),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ca, cb = rm.cell(row=r, column=1, value=a), rm.cell(row=r, column=2, value=b)
        ca.font = Font(name=F, bold=True, size=14 if r == 1 else 10)
        cb.font = body
        ca.alignment = cb.alignment = wrap
    for ref in ("B5", "B6"):
        rm[ref].font = Font(name=F, bold=True, size=12)
        rm[ref].alignment = Alignment(horizontal="left")
    rm.column_dimensions["A"].width = 22
    rm.column_dimensions["B"].width = 100
    leg = rm.cell(row=len(lines) + 2, column=1, value="Yellow cells")
    leg.font, leg.fill = Font(name=F, bold=True, size=10), fill_in
    rm.cell(row=len(lines) + 2, column=2, value="The only cells to fill in (Strings tab, columns H and I).").font = body
    ex = len(lines) + 3
    rm.cell(row=ex, column=1, value="Example").font = Font(name=F, bold=True, size=10)
    c = rm.cell(row=ex, column=2, value=cfg["example"])
    c.font, c.alignment = body, wrap
    wb.save(out)
    print(f"{out}\nrows={len(rows)} words={total}")



if __name__ == "__main__":
    ap = argparse.ArgumentParser(usage=__doc__)
    ap.add_argument("lang", nargs="?", choices=sorted(LANGS))
    ap.add_argument("out", nargs="?")
    ap.add_argument("--root", default=".")
    ap.add_argument("--store", nargs=2, metavar=("SRC", "DRAFT"))
    ap.add_argument("--texts", metavar="DIR")
    ap.add_argument("--store-src")
    ap.add_argument("--store-dir")
    a = ap.parse_args()
    if bool(a.store_src) != bool(a.store_dir):
        ap.error("--store-src and --store-dir go together")
    if a.texts:
        write_texts(a.texts, a.root, a.store_src, a.store_dir)
    elif a.lang and a.out:
        build(a.lang, a.out, a.root, a.store)
    else:
        ap.error("give <lang> <out.xlsx>, or --texts DIR")
